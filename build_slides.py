#!/usr/bin/env python
"""Build MolParetoBO presentation slides (PPTX + PDF)."""
from __future__ import annotations

import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
FIG = ROOT / "figures"
OUT_PPTX = ROOT / "MolParetoBO_Slides.pptx"
OUT_PDF = ROOT / "MolParetoBO_Slides.pdf"
OUT_HTML = ROOT / "MolParetoBO_Slides.html"

ACCENT = RGBColor(0x0A, 0x27, 0x48)
INK = RGBColor(0x15, 0x20, 0x2B)
MUTED = RGBColor(0x5A, 0x6A, 0x78)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_run(run, text, size=20, bold=False, color=INK, font="Calibri"):
  run.text = text
  run.font.size = Pt(size)
  run.font.bold = bold
  run.font.color.rgb = color
  run.font.name = font


def add_bg(slide, prs, color=WHITE):
  shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
  shape.fill.solid()
  shape.fill.fore_color.rgb = color
  shape.line.fill.background()
  # send to back
  spTree = slide.shapes._spTree
  sp = shape._element
  spTree.remove(sp)
  spTree.insert(2, sp)


def add_accent_bar(slide, prs):
  bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.18), prs.slide_height)
  bar.fill.solid()
  bar.fill.fore_color.rgb = ACCENT
  bar.line.fill.background()


def title_slide(prs, title, subtitle, lines):
  slide = prs.slides.add_slide(prs.slide_layouts[6])
  add_bg(slide, prs)
  band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.15))
  band.fill.solid()
  band.fill.fore_color.rgb = ACCENT
  band.line.fill.background()

  box = slide.shapes.add_textbox(Inches(0.7), Inches(1.6), Inches(12), Inches(1.2))
  tf = box.text_frame
  p = tf.paragraphs[0]
  run = p.add_run()
  set_run(run, title, size=40, bold=True, color=INK)

  box2 = slide.shapes.add_textbox(Inches(0.7), Inches(2.8), Inches(11.5), Inches(1.0))
  tf2 = box2.text_frame
  tf2.word_wrap = True
  p2 = tf2.paragraphs[0]
  r2 = p2.add_run()
  set_run(r2, subtitle, size=20, color=MUTED)

  y = 4.0
  for line in lines:
    box = slide.shapes.add_textbox(Inches(0.7), Inches(y), Inches(11), Inches(0.4))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    set_run(r, line, size=16, color=INK)
    y += 0.38
  return slide


def content_slide(prs, title, bullets, figure=None, note=None):
  slide = prs.slides.add_slide(prs.slide_layouts[6])
  add_bg(slide, prs)
  add_accent_bar(slide, prs)

  tbox = slide.shapes.add_textbox(Inches(0.55), Inches(0.35), Inches(12), Inches(0.7))
  tr = tbox.text_frame.paragraphs[0].add_run()
  set_run(tr, title, size=28, bold=True, color=INK)

  width = Inches(6.2) if figure else Inches(12.0)
  box = slide.shapes.add_textbox(Inches(0.55), Inches(1.15), width, Inches(5.5))
  tf = box.text_frame
  tf.word_wrap = True
  for i, b in enumerate(bullets):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.level = 0
    p.space_after = Pt(10)
    r = p.add_run()
    set_run(r, "• " + b, size=18, color=INK)

  if figure and Path(figure).exists():
    slide.shapes.add_picture(str(figure), Inches(7.1), Inches(1.3), width=Inches(5.5))

  if note:
    nbox = slide.shapes.add_textbox(Inches(0.55), Inches(6.85), Inches(12), Inches(0.4))
    nr = nbox.text_frame.paragraphs[0].add_run()
    set_run(nr, note, size=12, color=MUTED)
  return slide


def build_pptx():
  prs = Presentation()
  prs.slide_width = Inches(13.333)
  prs.slide_height = Inches(7.5)

  title_slide(
    prs,
    "MolParetoBO",
    "Surrogate-assisted multi-objective molecular optimization",
    [
      "Master in AI for Drug Discovery — University of Salerno",
      "Student: Ali Ghiami · Project 06",
      "10–15 minute presentation",
    ],
  )

  content_slide(
    prs,
    "1. Why multi-objective?",
    [
      "Drug candidates must balance many conflicting properties",
      "A single weighted score hides trade-offs",
      "Pareto optimization keeps undominated compromises",
      "Expensive oracles cannot evaluate every mutant",
      "Need: mutations + filters + surrogate + smart selection",
    ],
  )

  content_slide(
    prs,
    "2. Pipeline at a glance",
    [
      "Seed → RDKit mutations → suitability filters",
      "Oracle scores 5 endpoints (QED, SA, logP, logS, Tox)",
      "Surrogate (RF on Morgan fingerprints) predicts cheaply",
      "Pareto front shows trade-offs",
      "Acquisition picks the next batch (quality + uncertainty + diversity)",
    ],
    figure=FIG / "05_optimization_progress.png",
  )

  content_slide(
    prs,
    "3. Seed & endpoints",
    [
      "Seed: ibuprofen (one molecule first, then extend)",
      "QED ↑ · SA ↓ · logP ~ 2.5 · logS ↑ · Tox ↓",
      "Oracles are computational proxies (student setting)",
      "Toxicity is a teaching heuristic — stated explicitly",
      "SoftMining-style workflow; public seed for reproducibility",
    ],
    figure=FIG / "01_ibuprofen_atom_indices.png",
  )

  content_slide(
    prs,
    "4. Mutations + suitability",
    [
      "Methyl add/remove and ring F/Cl/OH/Me",
      "After mutation: sanitize, deduplicate, MW, QED gates",
      "kept column = suitability decision",
      "Gen-2 mutants enlarge surrogate training data",
      "Interactive widgets: tune QED filter and re-run live",
    ],
    figure=FIG / "02_mutation_gallery.png",
  )

  content_slide(
    prs,
    "5. Surrogate model",
    [
      "Morgan fingerprints (radius 2, 2048 bits)",
      "Multi-output Random Forest for 5 endpoints",
      "Hold-out MAE / R² reported honestly",
      "Small chemical neighbourhood → limited generalization",
      "Retrained inside the active loop as labels arrive",
    ],
    figure=FIG / "03_surrogate_parity.png",
  )

  content_slide(
    prs,
    "6. Pareto front",
    [
      "Maximize: +QED, −SA, −|logP−2.5|, +logS, −Tox",
      "Non-dominated molecules = approximate front",
      "No single ‘best’ compound — trade-offs are the result",
      "2D projections for interpretation",
      "Hypervolume proxy monitors progress",
    ],
    figure=FIG / "04_pareto_front.png",
  )

  content_slide(
    prs,
    "7. Acquisition & active loop",
    [
      "acquisition = α·quality + β·uncertainty + γ·diversity",
      "Select top-k → oracle label → retrain → update front",
      "5 iterations demonstrated in the notebook",
      "Widget lets you change α/β/γ live",
      "Simplified alternative to full qEHVI (proposal-friendly)",
    ],
    figure=FIG / "05_optimization_progress.png",
  )

  content_slide(
    prs,
    "8. Discussion — strengths",
    [
      "End-to-end reproducible workflow (local + Colab)",
      "Clear scientific framing of multi-objective trade-offs",
      "AI components: representations, learning, active selection",
      "Interactive demos verification",
      "Honest limits stated for evaluation criteria",
    ],
  )

  content_slide(
    prs,
    "9. Discussion — limits",
    [
      "Computational oracles ≠ wet-lab assays",
      "Tox proxy is pedagogical, not validated toxicology",
      "Local chemistry around ibuprofen",
      "RF tree variance ≠ calibrated Bayesian uncertainty",
      "2D HV monitor understates full 5-D volume",
    ],
  )

  content_slide(
    prs,
    "10. Extensions — GP",
    [
      "PCA on fingerprints → Gaussian Process",
      "Endpoint-wise predictive standard deviation",
      "Drop-in uncertainty for acquisition",
      "Scales better than raw 2048-D GP on small N",
      "Runnable in the notebook (Section 10A)",
    ],
    figure=FIG / "06_gp_uncertainty.png" if (FIG / "06_gp_uncertainty.png").exists() else None,
  )

  content_slide(
    prs,
    "10b. Tools & interactive widgets",
    [
      "RDKit · scikit-learn · pymoo · NumPy/pandas/matplotlib",
      "ipywidgets Control Panel: 5 one-click checks",
      "1 filters · 2 inspect · 3 surrogate · 4 acquisition · 5 Pareto",
      "Runs on local Jupyter or Google Colab",
      "Details: Appendix B in the notebook / report",
    ],
    note="Open notebook → Run all → Section 8c → click buttons 1–5",
  )

  content_slide(
    prs,
    "11. Extensions — NSGA-II & EHVI",
    [
      "NSGA-II: non-dominated sorting + crowding on mutant pool",
      "Keeps a diverse elite subset of the front",
      "2D EHVI: Monte Carlo expected hypervolume improvement",
      "Suggests candidates that grow the front",
      "Full qEHVI (BoTorch) listed as future work",
    ],
    figure=FIG / "08_ehvi_picks.png" if (FIG / "08_ehvi_picks.png").exists() else (
      FIG / "07_nsga_subset.png" if (FIG / "07_nsga_subset.png").exists() else None
    ),
  )

  content_slide(
    prs,
    "12. Conclusions",
    [
      "Built a teaching-scale MolParetoBO loop",
      "Mutate → filter → score → surrogate → Pareto → acquire",
      "Results are trade-off sets, not one winner",
      "Extensions show path to research-grade BO",
      "Ready for SoftMining / tutor seeds next",
    ],
    note="Questions? · Notebook: MolParetoBO_Report.ipynb · Report: MolParetoBO_Report.pdf",
  )

  prs.save(OUT_PPTX)
  print("Wrote", OUT_PPTX)


def build_html_pdf():
  """Also emit a simple HTML slide deck and print to PDF via Chrome."""
  import base64
  import pandas as pd

  def b64(p: Path) -> str:
    if not p.exists():
      return ""
    return base64.b64encode(p.read_bytes()).decode()

  def fig(name, caption):
    p = FIG / name
    if not p.exists():
      return f"<p class='muted'>{caption} (figure pending)</p>"
    return f"<img src='data:image/png;base64,{b64(p)}' alt='{caption}'/><p class='cap'>{caption}</p>"

  hist = ""
  hp = ROOT / "data" / "optimization_history.csv"
  if hp.exists():
    df = pd.read_csv(hp)
    hist = df.to_html(index=False, classes="t")

  slides = [
    ("MolParetoBO", "<p class='sub'>Surrogate-assisted multi-objective molecular optimization</p><p>Ali Ghiami · University of Salerno · AI for Drug Discovery · Project 06</p>"),
    ("Why multi-objective?", "<ul><li>Conflicting ADMET / drug-likeness goals</li><li>Weighted scores hide trade-offs</li><li>Pareto keeps undominated compromises</li><li>Oracles are expensive → surrogates + acquisition</li></ul>"),
    ("Pipeline", "<ol><li>Mutate (RDKit)</li><li>Suitability filters</li><li>Score 5 endpoints</li><li>Surrogate (RF)</li><li>Pareto + acquisition loop</li></ol>" + fig("05_optimization_progress.png", "Active loop progress")),
    ("Seed & endpoints", "<p><b>Ibuprofen</b> + QED↑ SA↓ logP~2.5 logS↑ Tox↓</p>" + fig("01_ibuprofen_atom_indices.png", "Seed with atom indices")),
    ("Mutations & filters", "<p>Methyl / ring edits · sanitize · MW · QED gate</p>" + fig("02_mutation_gallery.png", "Mutation gallery")),
    ("Surrogate", fig("03_surrogate_parity.png", "Surrogate vs oracle")),
    ("Pareto front", fig("04_pareto_front.png", "Pareto projections")),
    ("Acquisition loop", "<p>α·quality + β·uncertainty + γ·diversity</p>" + (hist or "")),
    ("Discussion", "<ul><li><b>Strength:</b> full reproducible AI workflow + widgets</li><li><b>Limit:</b> proxy oracles, local chemistry, RF uncertainty</li><li><b>Meaning:</b> front = trade-offs, not one best drug</li></ul>"),
    ("Extensions", "<ul><li>GP + PCA uncertainty</li><li>NSGA-II crowding subset</li><li>2D EHVI → roadmap to qEHVI/BoTorch</li></ul>" + fig("08_ehvi_picks.png", "EHVI picks") + fig("06_gp_uncertainty.png", "GP uncertainty")),
    ("Tools & widgets", "<ul><li><b>RDKit</b> mutations/descriptors · <b>sklearn</b> RF/GP · <b>pymoo</b> NSGA-II</li><li><b>ipywidgets</b> Interactive Control Panel (5 checks)</li><li>1 filters · 2 inspect · 3 surrogate · 4 acquisition · 5 Pareto</li><li>Local Jupyter or Google Colab</li></ul>"),
    ("Conclusions", "<ul><li>Teaching-scale MolParetoBO delivered</li><li>Colab-ready · report · slides · widgets</li><li>Ready for SoftMining seeds / stronger oracles</li></ul><p class='sub'>Thank you — questions?</p>"),
  ]

  parts = []
  for i, (title, body) in enumerate(slides, 1):
    parts.append(f"<section class='slide'><div class='num'>{i}/{len(slides)}</div><h1>{title}</h1>{body}</section>")

  html = f"""<!DOCTYPE html><html><head><meta charset='utf-8'/>
<title>MolParetoBO Slides</title>
<style>
@import url('https://fonts.googleapis.com/css2?family=Source+Sans+3:wght@400;600;700&family=Source+Serif+4:opsz,wght@8..60,600&display=swap');
html,body{{margin:0;padding:0;background:#fff;color:#15202b;font-family:'Source Sans 3',sans-serif;}}
.slide{{min-height:100vh;padding:48px 64px;page-break-after:always;border-left:10px solid #0a2748;box-sizing:border-box;}}
h1{{font-family:'Source Serif 4',Georgia,serif;font-size:2.2rem;margin:0 0 18px;}}
.sub{{color:#5a6a78;font-size:1.15rem;}}
ul,ol{{font-size:1.25rem;line-height:1.45;max-width:48ch;}}
img{{max-width:min(720px,90%);border:1px solid #d5dee6;border-radius:10px;margin-top:12px;}}
.cap{{color:#5a6a78;font-size:.9rem;}}
.num{{color:#0a2748;font-weight:700;letter-spacing:.08em;font-size:.8rem;}}
table.t{{border-collapse:collapse;font-size:.85rem;margin-top:12px;}}
table.t th,table.t td{{border-bottom:1px solid #d5dee6;padding:6px 8px;text-align:left;}}
</style></head><body>{''.join(parts)}</body></html>"""
  OUT_HTML.write_text(html, encoding="utf-8")
  chrome = Path("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome")
  subprocess.run(
    [str(chrome), "--headless", "--disable-gpu", "--no-pdf-header-footer",
     f"--print-to-pdf={OUT_PDF}", OUT_HTML.as_uri()],
    check=True, capture_output=True,
  )
  print("Wrote", OUT_HTML)
  print("Wrote", OUT_PDF, OUT_PDF.stat().st_size)


if __name__ == "__main__":
  build_pptx()
  build_html_pdf()
